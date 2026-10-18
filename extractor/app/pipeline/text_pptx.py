"""Stage 1 for PPTX: slide title, body paragraphs, tables, and speaker notes."""

from __future__ import annotations

import contextlib
from collections.abc import Iterator
from pathlib import Path

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE, PP_PLACEHOLDER
from pptx.exc import PackageNotFoundError
from pptx.shapes.base import BaseShape
from pptx.slide import Slide

from .errors import CorruptFileError, TooManyPagesError
from .ocr import needs_ocr, ocr_image_bytes
from .types import Section

_TITLE_TYPES = {PP_PLACEHOLDER.TITLE, PP_PLACEHOLDER.CENTER_TITLE, PP_PLACEHOLDER.VERTICAL_TITLE}
MAX_OCR_IMAGES_PER_SLIDE = 5


def _walk(shapes: object) -> Iterator[BaseShape]:
    for shape in shapes:  # type: ignore[attr-defined]
        if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
            yield from _walk(shape.shapes)
        else:
            yield shape


def _is_title(shape: BaseShape) -> bool:
    if not shape.is_placeholder:
        return False
    try:
        return shape.placeholder_format.type in _TITLE_TYPES
    except ValueError:
        return False


def _shape_lines(shape: BaseShape) -> list[str]:
    lines: list[str] = []
    if shape.has_text_frame:
        for para in shape.text_frame.paragraphs:  # type: ignore[attr-defined]
            text = "".join(run.text for run in para.runs).strip() or para.text.strip()
            if not text:
                continue
            # Body placeholders render every paragraph as a bullet.
            is_bullet = shape.is_placeholder or para.level > 0
            lines.append(f"- {text}" if is_bullet else text)
        lines.append("")
    if getattr(shape, "has_table", False) and shape.has_table:
        for row in shape.table.rows:  # type: ignore[attr-defined]
            cells = [c.text.strip().replace("\n", " ") for c in row.cells]
            cells = [c for c in cells if c]
            if len(cells) == 2:
                lines.append(f"{cells[0]}: {cells[1]}")
            elif cells:
                lines.append(" | ".join(cells))
        lines.append("")
    return lines


def _slide_section(slide: Slide, number: int, ocr_min_chars: int, ocr_language: str) -> Section:
    title: str | None = None
    body: list[str] = []
    pictures: list[bytes] = []
    for shape in _walk(slide.shapes):
        if _is_title(shape) and shape.has_text_frame and title is None:
            title = shape.text_frame.text.strip() or None  # type: ignore[attr-defined]
            continue
        body.extend(_shape_lines(shape))
        if shape.shape_type == MSO_SHAPE_TYPE.PICTURE and len(pictures) < MAX_OCR_IMAGES_PER_SLIDE:
            with contextlib.suppress(AttributeError, KeyError, ValueError):
                pictures.append(shape.image.blob)  # type: ignore[attr-defined]
    notes: list[str] = []
    if slide.has_notes_slide and slide.notes_slide.notes_text_frame is not None:
        notes = slide.notes_slide.notes_text_frame.text.splitlines()

    ocr = False
    if pictures and needs_ocr("\n".join(body), ocr_min_chars):
        ocr_lines: list[str] = []
        for blob in pictures:
            ocr_lines.extend(ocr_image_bytes(blob, ocr_language).splitlines())
            ocr_lines.append("")
        if any(line.strip() for line in ocr_lines):
            body.extend(ocr_lines)
            ocr = True
    return Section(
        ref=f"slide {number}", lines=body, title=title, notes=notes, ocr=ocr, is_slide=True
    )


def extract_pptx(
    path: Path, *, max_pages: int, ocr_min_chars: int, ocr_language: str = "eng"
) -> list[Section]:
    try:
        prs = Presentation(str(path))
    except (PackageNotFoundError, KeyError, ValueError) as exc:
        raise CorruptFileError() from exc
    slides = list(prs.slides)
    if len(slides) > max_pages:
        raise TooManyPagesError(f"The file has {len(slides)} slides; the limit is {max_pages}.")
    return [
        _slide_section(slide, i + 1, ocr_min_chars, ocr_language) for i, slide in enumerate(slides)
    ]
