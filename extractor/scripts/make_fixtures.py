"""Generate the test fixtures in tests/fixtures/.

Run from extractor/:  python scripts/make_fixtures.py

The generated files are committed so tests do not depend on this script, but
it documents exactly what each fixture contains and lets anyone rebuild them.
Every detector is covered by at least one fixture:

  qa_worksheet.pdf     labeled pairs, numbered/lettered questions, multiple
                       choice (inline answer and answer key), fill in the
                       blank, repeated header/footer, page numbers, hyphenation
  glossary.pdf         term: definition, term - definition, definitional
                       sentences
  lecture.pptx         slide title + bullets, title question + notes answer,
                       two-column table, multiple choice slide, generic
                       "Agenda" slide that must be skipped
  scanned.pdf          text rendered to an image and wrapped in a PDF with no
                       text layer (OCR path)
  image_slides.pptx    slide whose only content is a picture of text (PPTX OCR)
  encrypted.pdf        password protected (must be rejected)
  macro.pptm           macro-enabled presentation (must be rejected)
  fake.pdf             plain text with a .pdf name (bad magic bytes)
"""

from __future__ import annotations

import io
import shutil
import zipfile
from pathlib import Path

import pymupdf
from pptx import Presentation
from pptx.util import Inches, Pt
from reportlab.lib.pagesizes import letter
from reportlab.pdfgen import canvas

OUT = Path(__file__).resolve().parent.parent / "tests" / "fixtures"

WORKSHEET_PAGES = [
    [
        "Q: What is the capital of France?",
        "A: Paris",
        "",
        "Question 2: Who wrote the play Hamlet?",
        "Answer: William Shakespeare",
        "",
        "3. What is the chemical symbol for gold?",
        "Au",
        "",
        "4. Why do we have seasons on Earth? Because the Earth's axis is tilted rel-",
        "ative to its orbit around the Sun.",
    ],
    [
        "5. Which planet is known as the Red Planet?",
        "A) Venus",
        "B) Mars",
        "C) Jupiter",
        "D) Saturn",
        "Answer: B",
        "",
        "6. Which gas do plants absorb from the air?",
        "A. Oxygen",
        "B. Nitrogen",
        "C. Carbon dioxide",
        "D. Helium",
        "",
        "a) What is the largest ocean on Earth?",
        "The Pacific Ocean",
    ],
    [
        "7. The powerhouse of the cell is the ____. (mitochondria)",
        "8. Water is made of hydrogen and [blank].",
        "",
        "Answer Key",
        "6. C",
        "8. oxygen",
    ],
]

GLOSSARY_LINES = [
    "Biology Glossary",
    "",
    "Osmosis: the movement of water across a semipermeable membrane",
    "Mitosis - cell division that produces two identical daughter cells",
    "Enzyme: a protein that speeds up a chemical reaction",
    "",
    "Photosynthesis is the process by which plants convert light energy into",
    "chemical energy.",
    "A gene refers to a sequence of DNA that codes for a protein.",
    "It is important to review these terms every week.",
]

SCANNED_LINES = [
    "Q: What is the speed of light?",
    "A: About 300,000 kilometers per second",
    "",
    "Gravity: the force that attracts objects toward each other",
]


def _pdf(path: Path, pages: list[list[str]], header: str | None = None) -> None:
    c = canvas.Canvas(str(path), pagesize=letter, invariant=1)
    c.setTitle("Quiz.me fixture")
    width, height = letter
    for number, lines in enumerate(pages, start=1):
        if header:
            c.setFont("Helvetica", 9)
            c.drawString(72, height - 40, header)
            c.drawCentredString(width / 2, 30, f"Page {number} of {len(pages)}")
        c.setFont("Helvetica", 12)
        y = height - 80
        for line in lines:
            c.drawString(72, y, line)
            y -= 18
        c.showPage()
    c.save()


def make_worksheet() -> None:
    _pdf(OUT / "qa_worksheet.pdf", WORKSHEET_PAGES, header="General Science Review - Unit 4")


def make_glossary() -> None:
    _pdf(OUT / "glossary.pdf", [GLOSSARY_LINES])


def _text_image_png(lines: list[str], scale: float = 3.0) -> bytes:
    """Render lines with PDF base-14 fonts, then rasterize: no system fonts needed."""
    doc = pymupdf.open()
    page = doc.new_page(width=612, height=40 + 22 * len(lines))
    y = 30.0
    for line in lines:
        page.insert_text((36, y), line, fontsize=14, fontname="helv")
        y += 22
    pix = page.get_pixmap(matrix=pymupdf.Matrix(scale, scale), colorspace=pymupdf.csGRAY)
    data = pix.tobytes("png")
    doc.close()
    return bytes(data)


def make_scanned() -> None:
    png = _text_image_png(SCANNED_LINES)
    doc = pymupdf.open()
    page = doc.new_page(width=612, height=792)
    page.insert_image(
        pymupdf.Rect(36, 36, 576, 36 + 540 * (40 + 22 * len(SCANNED_LINES)) / 612), stream=png
    )
    doc.save(OUT / "scanned.pdf", garbage=4, deflate=True, no_new_id=True)
    doc.close()


def make_lecture() -> None:
    prs = Presentation()
    title_only = prs.slide_layouts[5]
    bullets = prs.slide_layouts[1]

    s = prs.slides.add_slide(bullets)
    s.shapes.title.text = "Agenda"
    s.placeholders[1].text_frame.text = "Cells"
    s.placeholders[1].text_frame.add_paragraph().text = "Energy"

    s = prs.slides.add_slide(bullets)
    s.shapes.title.text = "Functions of the cell membrane"
    tf = s.placeholders[1].text_frame
    tf.text = "Controls what enters and leaves the cell"
    tf.add_paragraph().text = "Protects the cell from its surroundings"

    s = prs.slides.add_slide(title_only)
    s.shapes.title.text = "What is ATP?"
    s.notes_slide.notes_text_frame.text = "The main energy currency of the cell"

    s = prs.slides.add_slide(title_only)
    s.shapes.title.text = "Key terms"
    table = s.shapes.add_table(3, 2, Inches(1), Inches(2), Inches(8), Inches(2)).table
    rows = [
        ("Term", "Meaning"),
        ("Ribosome", "the site of protein synthesis"),
        ("Nucleus", "the organelle that holds the genetic material"),
    ]
    for r, (a, b) in enumerate(rows):
        table.cell(r, 0).text = a
        table.cell(r, 1).text = b

    s = prs.slides.add_slide(title_only)
    s.shapes.title.text = "Quick check"
    box = s.shapes.add_textbox(Inches(1), Inches(2), Inches(8), Inches(3)).text_frame
    lines = [
        "1. Which organelle contains chlorophyll?",
        "A) Mitochondrion",
        "B) Chloroplast",
        "C) Ribosome",
        "Answer: B",
    ]
    box.text = lines[0]
    for line in lines[1:]:
        box.add_paragraph().text = line
    for p in box.paragraphs:
        p.font.size = Pt(20)

    prs.save(str(OUT / "lecture.pptx"))


def make_image_slides() -> None:
    prs = Presentation()
    s = prs.slides.add_slide(prs.slide_layouts[6])  # blank
    png = _text_image_png(["Q: What is the largest planet?", "A: Jupiter"])
    s.shapes.add_picture(io.BytesIO(png), Inches(0.5), Inches(1), width=Inches(9))
    prs.save(str(OUT / "image_slides.pptx"))


def make_encrypted() -> None:
    doc = pymupdf.open(OUT / "glossary.pdf")
    doc.save(
        OUT / "encrypted.pdf",
        encryption=pymupdf.PDF_ENCRYPT_AES_256,  # type: ignore[attr-defined]
        owner_pw="fixture-owner",
        user_pw="fixture-user",
    )
    doc.close()


def make_macro() -> None:
    src = OUT / "lecture.pptx"
    dst = OUT / "macro.pptm"
    with zipfile.ZipFile(src) as zin, zipfile.ZipFile(dst, "w", zipfile.ZIP_DEFLATED) as zout:
        for item in zin.infolist():
            data = zin.read(item.filename)
            if item.filename == "[Content_Types].xml":
                data = data.replace(
                    b"application/vnd.openxmlformats-officedocument.presentationml.presentation.main+xml",
                    b"application/vnd.ms-powerpoint.presentation.macroEnabled.main+xml",
                )
            zout.writestr(item, data)
        zout.writestr("ppt/vbaProject.bin", b"\x00" * 64)


def make_fake() -> None:
    (OUT / "fake.pdf").write_text("This is plain text pretending to be a PDF.\n")


def main() -> None:
    OUT.mkdir(parents=True, exist_ok=True)
    make_worksheet()
    make_glossary()
    make_scanned()
    make_lecture()
    make_image_slides()
    make_encrypted()
    make_macro()
    make_fake()
    shutil.rmtree(OUT / "__pycache__", ignore_errors=True)
    for f in sorted(OUT.iterdir()):
        print(f"{f.name:20} {f.stat().st_size:>8} bytes")


if __name__ == "__main__":
    main()
